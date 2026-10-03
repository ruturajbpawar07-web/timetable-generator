"""Timetable / report exports as CSV or Excel.  (PDF = browser print of the timetable page.)"""
import csv
import io

from sqlalchemy import select

from . import db, services as svc
from scheduler import validator


def _grid(p, es, keep, title):
    """Weekly grid: rows = periods, columns = days."""
    head = ["Period"] + [d.name for d in p.days]
    rows = []
    for s in p.slots:
        if s.is_break:
            rows.append([f"{s.start}-{s.end}"] + ["BREAK"] * len(p.days))
            continue
        row = [f"{s.start}-{s.end}"]
        for d in p.days:
            cell = [f"{p.S[e.subject_id].code} {e.component} | {p.F[e.faculty_id].name} | {p.R[e.room_id].code} | "
                    f"{validator.cls_name(p, e)}" for e in es if keep(e) and e.day_id == d.id
                    and p.pos[s.id] in (validator.span(p, e) or [])]
            row.append("\n".join(cell))
        rows.append(row)
    return title, head, rows


def build(s, tid, kind, fmt, oid=None):
    p, es = svc.problem_and_entries(s, tid)
    if kind == "class":
        title, head, rows = _grid(p, es, lambda e: e.class_id == oid, f"Class {p.C[oid].name}")
    elif kind == "faculty":
        title, head, rows = _grid(p, es, lambda e: e.faculty_id == oid, p.F[oid].name)
    elif kind == "room":
        title, head, rows = _grid(p, es, lambda e: e.room_id == oid, f"Room {p.R[oid].code}")
    elif kind == "master":
        title, head = "Master timetable", ["Day", "Start", "End", "Class", "Batch", "Subject", "Type", "Faculty", "Room", "Locked"]
        rows = [[p.DAY[e.day_id].name, p.slots[p.pos[e.slot_id]].start,
                 p.slots[p.pos[e.slot_id] + e.duration - 1].end, p.C[e.class_id].name,
                 next((b.name for b in p.C[e.class_id].batches if b.id == e.batch_id), ""), p.S[e.subject_id].name,
                 e.component, p.F[e.faculty_id].name, p.R[e.room_id].code, "yes" if e.locked else ""]
                for e in sorted(es, key=lambda e: (e.day_id, p.pos[e.slot_id], p.C[e.class_id].name))]
    elif kind == "workload":
        w = svc.workload(s, tid)
        title = "Faculty workload"
        head = ["Faculty", "Department", *[d.name for d in p.days], "Weekly", "Min", "Max", "Utilization %", "Status"]
        rows = [[r["name"], r["department"], *r["daily"].values(), r["weekly"], r["min_per_week"], r["max_per_week"],
                 r["utilization"], r["status"]] for r in w["faculty"]]
    elif kind == "substitutions":
        title = "Substitution report"
        head = ["Date", "Status", "Class", "Subject", "Slot", "Original", "Substitute", "Reason"]
        by_id = {e.id: e for e in es}
        rows = []
        for x in s.scalars(select(db.Substitution).order_by(db.Substitution.date)):
            e = by_id.get(x.entry_id)
            if e:
                rows.append([x.date.isoformat(), x.status, validator.cls_name(p, e), p.S[e.subject_id].name,
                             validator.when(p, e.day_id, e.slot_id, e.duration), p.F[x.original_faculty_id].name,
                             p.F[x.substitute_faculty_id].name if x.substitute_faculty_id else "",
                             "; ".join(x.reasons or []) or (x.reason or "")])
    else:
        raise ValueError(f"unknown export kind {kind}")

    fname = f"{title.lower().replace(' ', '_').replace('.', '')}"
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(head)
        w.writerows(rows)
        return buf.getvalue().encode(), "text/csv", f"{fname}.csv"
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[1].column_letter].width = 28 if kind in ("class", "faculty", "room") else 18
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{fname}.xlsx"
