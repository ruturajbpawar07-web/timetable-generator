"""Load a college dataset from a folder of CSV or XLSX files into the database.

    python -m backend.importer data/dummy        # or data/college for real data

Each table is read from <name>.xlsx if present, else <name>.csv (first sheet,
first row = header).  Column names are documented in docs/DATA_IMPORT.md.
Rows reference each other by *codes* (F01, CE301, SE-CE-A); the importer
resolves them to foreign keys and fails loudly on unknown codes.
"""
import csv
import sys
from pathlib import Path

from . import db


def read(folder: Path, name: str) -> list[dict]:
    x, c = folder / f"{name}.xlsx", folder / f"{name}.csv"
    if x.exists():
        from openpyxl import load_workbook
        rows = list(load_workbook(x, read_only=True, data_only=True).active.iter_rows(values_only=True))
        head = [str(h).strip() for h in rows[0]]
        return [{h: ("" if v is None else str(v).strip()) for h, v in zip(head, r)} for r in rows[1:] if any(r)]
    if c.exists():
        with open(c, newline="", encoding="utf-8-sig") as f:
            return [{k.strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]
    raise FileNotFoundError(f"Missing {name}.csv / {name}.xlsx in {folder}")


def ref(table: dict, code: str, what: str, where: str):
    if code not in table:
        raise ValueError(f"{where}: unknown {what} '{code}'")
    return table[code]


def truthy(v: str) -> bool:
    return v.lower() in ("1", "true", "yes", "y")


def load(folder: str | Path):
    folder = Path(folder)
    db.Base.metadata.drop_all(db.engine)
    db.init()
    with db.Session() as s:
        settings = {r["key"]: r for r in read(folder, "settings")}
        for r in settings.values():
            s.add(db.ConstraintConfiguration(key=r["key"], value=r["value"], description=r.get("description")))
        college = db.College(name=settings.get("college_name", {}).get("value", "College"))
        ay = db.AcademicYear(label=settings.get("academic_year", {}).get("value", ""))
        s.add_all([college, ay])
        s.flush()

        dept, prog = {}, {}
        for r in read(folder, "departments"):
            d = dept[r["code"]] = db.Department(college_id=college.id, code=r["code"], name=r["name"])
            s.add(d); s.flush()
            prog[r["program_code"]] = db.Program(department_id=d.id, code=r["program_code"], name=r["program_name"])
            s.add(prog[r["program_code"]])

        days = {}
        for r in read(folder, "working_days"):
            days[r["name"]] = db.WorkingDay(name=r["name"], order=int(r["order"]))
        slots = {}
        for r in read(folder, "time_slots"):
            slots[r["index"]] = db.TimeSlot(index=int(r["index"]), start=r["start"], end=r["end"],
                                            is_break=truthy(r["is_break"]), label=r.get("label"))
        rooms = {}
        for i, r in enumerate(read(folder, "rooms"), 2):
            rooms[r["code"]] = db.Room(code=r["code"], name=r["name"], kind=r["kind"].upper(), capacity=int(r["capacity"]),
                                       lab_type=r.get("lab_type") or None,
                                       department_id=ref(dept, r["department"], "department", f"rooms row {i}").id
                                       if r.get("department") else None)
        s.add_all([*days.values(), *slots.values(), *rooms.values()])
        s.flush()

        subj = {}
        for i, r in enumerate(read(folder, "subjects"), 2):
            subj[r["code"]] = db.Subject(
                code=r["code"], name=r["name"], department_id=ref(dept, r["department"], "department", f"subjects row {i}").id,
                semester_no=int(r["semester"] or 0), credits=float(r["credits"]),
                lecture_sessions_per_week=int(r["lecture_sessions_per_week"] or 0),
                tutorial_sessions_per_week=int(r["tutorial_sessions_per_week"] or 0),
                practical_sessions_per_week=int(r["practical_sessions_per_week"] or 0),
                practical_duration_periods=int(r["practical_duration_periods"] or 2),
                lab_type=r.get("lab_type") or None, is_core=truthy(r.get("is_core", "")),
                morning_preferred=truthy(r.get("morning_preferred", "")))
        s.add_all(subj.values())

        sems, div, batch = {}, {}, {}
        for i, r in enumerate(read(folder, "classes"), 2):
            n = int(r["semester"])
            if n not in sems:
                sems[n] = db.Semester(academic_year_id=ay.id, number=n)
                s.add(sems[n]); s.flush()
            d = div[r["code"]] = db.Division(
                code=r["code"], program_id=ref(prog, r["program"], "program", f"classes row {i}").id,
                semester_id=sems[n].id, year_level=r["year_level"], name=r["division"], strength=int(r["strength"]),
                home_room_id=ref(rooms, r["home_room"], "room", f"classes row {i}").id if r.get("home_room") else None)
            s.add(d); s.flush()
            for b in filter(None, r.get("batches", "").split(";")):
                name, strength = b.split(":")
                batch[r["code"], name] = db.StudentBatch(division_id=d.id, name=name, strength=int(strength))
                s.add(batch[r["code"], name])

        fac = {}
        for i, r in enumerate(read(folder, "faculty"), 2):
            fac[r["code"]] = db.Faculty(
                code=r["code"], name=r["name"], department_id=ref(dept, r["department"], "department", f"faculty row {i}").id,
                designation=r.get("designation"), max_per_day=int(r["max_per_day"]), max_per_week=int(r["max_per_week"]),
                min_per_week=int(r.get("min_per_week") or 0), max_consecutive=int(r.get("max_consecutive") or 3),
                subjects=[ref(subj, c, "subject", f"faculty row {i}") for c in r["subjects"].split(";") if c])
        s.add_all(fac.values())
        s.flush()

        for i, r in enumerate(read(folder, "assignments"), 2):
            where = f"assignments row {i}"
            d, su = ref(div, r["class"], "class", where), ref(subj, r["subject"], "subject", where)
            comp = r["component"].upper()
            default = {"L": (su.lecture_sessions_per_week, 1), "T": (su.tutorial_sessions_per_week, 1),
                       "P": (su.practical_sessions_per_week, su.practical_duration_periods)}[comp]
            req = db.SubjectRequirement(
                division_id=d.id, subject_id=su.id, component=comp,
                batch_id=ref(batch, (r["class"], r["batch"]), "batch", where).id if r.get("batch") else None,
                sessions_per_week=int(r.get("sessions") or default[0]), duration_periods=int(r.get("duration") or default[1]))
            s.add(req); s.flush()
            s.add(db.FacultyClassAssignment(requirement_id=req.id,
                                            faculty_id=ref(fac, r["faculty"], "faculty", where).id if r.get("faculty") else None))

        for i, r in enumerate(read(folder, "availability"), 2):
            where = f"availability row {i}"
            s.add(db.FacultyAvailability(faculty_id=ref(fac, r["faculty"], "faculty", where).id,
                                         day_id=ref(days, r["day"], "day", where).id,
                                         slot_id=ref(slots, r["period_index"], "period", where).id, kind=r["kind"].upper()))
        s.add(db.AuditLog(action="DATA_IMPORT", new={"folder": str(folder)}, reason="Dataset imported"))
        s.commit()


if __name__ == "__main__":
    load(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent.parent / "data" / "dummy")
    print("imported")
